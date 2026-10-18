"""Downloadable test reports for one run: an Excel workbook with every detail, and a standalone HTML page.

Excel sheets: Summary, Differences, Root causes, All cases, Failures (message + stack trace per failing
test), and the build log of each side. The HTML report has the same content on one page that opens in any
browser, prints cleanly and can be attached to an email.
"""

from __future__ import annotations

import html
import io
from datetime import datetime

import pandas as pd

from .compare import (
    CHANGED, FIXED, ONLY_LEGACY, ONLY_NEWGEN, OUTPUT_DIFF, REGRESSION, SAME, Comparison, category_counts,
    difference_causes,
)
from .java_runner import RunResult, TestCase

SIDE_NAMES = {"legacy": "Legacy", "newgen": "New gen"}
# Plain category names (no emoji) so reports read the same in Excel, browsers and print.
CATEGORY_NAMES = {
    REGRESSION: "Regression", FIXED: "Fixed", CHANGED: "Changed", OUTPUT_DIFF: "Output differs",
    ONLY_LEGACY: "Only in legacy", ONLY_NEWGEN: "Only in new gen", SAME: "Same",
}
CATEGORY_COLORS = {  # light fills for Excel and HTML
    REGRESSION: "FDE2E2", FIXED: "DBEAFE", CHANGED: "FEF3C7", OUTPUT_DIFF: "CCFBF1",
    ONLY_LEGACY: "E0F2FE", ONLY_NEWGEN: "EDE9FE", SAME: "FFFFFF",
}
EXCEL_CELL_LIMIT = 32_000  # Excel refuses cells over 32,767 characters


def _test_name(tc: TestCase | None) -> str:
    return f"{tc.suite.rsplit('.', 1)[-1]}.{tc.name}" if tc else ""


def cases_frame(rows: list[Comparison]) -> pd.DataFrame:
    """One row per input case with both sides side by side."""
    return pd.DataFrame([{
        "Case": (r.legacy or r.newgen).case_id or r.test_id,
        "Category": CATEGORY_NAMES[r.category],
        "Description": r.display_name,
        "Legacy status": r.legacy_status,
        "New gen status": r.newgen_status,
        "Legacy output": r.legacy_output,
        "New gen output": r.newgen_output,
        "Legacy message": r.legacy.message if r.legacy else "",
        "New gen message": r.newgen.message if r.newgen else "",
        "Legacy time (s)": round(r.legacy.time, 3) if r.legacy else None,
        "New gen time (s)": round(r.newgen.time, 3) if r.newgen else None,
        "Compare only (bulk)": "yes" if (r.legacy or r.newgen).compare_only else "",
        "Legacy test": _test_name(r.legacy),
        "New gen test": _test_name(r.newgen),
    } for r in rows], columns=["Case", "Category", "Description", "Legacy status", "New gen status",
                               "Legacy output", "New gen output", "Legacy message", "New gen message",
                               "Legacy time (s)", "New gen time (s)", "Compare only (bulk)", "Legacy test",
                               "New gen test"])


def causes_frame(rows: list[Comparison]) -> pd.DataFrame:
    return pd.DataFrame([{
        "Output word": c.token,
        "Only in": c.side,
        "Cases": len(c.cases),
        "Case ids": ", ".join((r.legacy or r.newgen).case_id or r.test_id for r in c.cases),
    } for c in difference_causes(rows)], columns=["Output word", "Only in", "Cases", "Case ids"])


def failures_frame(results: dict[str, RunResult]) -> pd.DataFrame:
    return pd.DataFrame([{
        "Side": SIDE_NAMES[side],
        "Case": t.case_id,
        "Test": _test_name(t),
        "Status": t.status,
        "Output": t.output,
        "Message": t.message,
        "Stack trace": t.details,
    } for side, r in results.items() for t in r.tests if t.status in ("failed", "error")],
        columns=["Side", "Case", "Test", "Status", "Output", "Message", "Stack trace"])


def summary_items(started: datetime, sources: dict[str, dict[str, str]], results: dict[str, RunResult],
                  rows: list[Comparison]) -> list[tuple[str, object]]:
    counts = category_counts(rows)
    items: list[tuple[str, object]] = [
        ("Run started", started.strftime("%Y-%m-%d %H:%M:%S")),
        ("Run time (s)", round(max(r.duration for r in results.values()), 1)),
        ("Input cases", len(rows)),
        ("Differ", len(rows) - counts[SAME]),
    ]
    items += [(CATEGORY_NAMES[k], v) for k, v in counts.items()]
    for side, r in results.items():
        name, c = SIDE_NAMES[side], r.counts()
        source = sources.get(side, {})
        items += [
            (f"{name} source", f"{source.get('origin', '')}: {source.get('detail', '')}".strip(": ")),
            (f"{name} build", "ran" if r.ok else f"stopped at {r.stage}: {r.error}"),
            (f"{name} passed", f"{c['passed']} of {c['total']}"),
            (f"{name} pass rate", f"{c['passed'] / c['total'] * 100:.1f}%" if c["total"] else "-"),
            (f"{name} failed / error / skipped", f"{c['failed']} / {c['error']} / {c['skipped']}"),
            (f"{name} bulk cases compared", c["bulk"]),
        ]
    return items


def _clip(value):
    if isinstance(value, str) and len(value) > EXCEL_CELL_LIMIT:
        return value[:EXCEL_CELL_LIMIT] + " …(truncated)"
    return value


def excel_report(started: datetime, sources: dict[str, dict[str, str]], results: dict[str, RunResult],
                 rows: list[Comparison]) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cases = cases_frame(rows)
    sheets = {
        "Summary": pd.DataFrame(summary_items(started, sources, results, rows), columns=["Item", "Value"]),
        "Differences": cases[cases["Category"] != CATEGORY_NAMES[SAME]],
        "Root causes": causes_frame(rows),
        "All cases": cases,
        "Failures": failures_frame(results),
    }
    for side, r in results.items():
        sheets[f"Build log {SIDE_NAMES[side]}"] = pd.DataFrame({"Line": (r.log or "(no output)").splitlines()})

    buffer = io.BytesIO()
    fills = {CATEGORY_NAMES[k]: PatternFill("solid", fgColor=v) for k, v in CATEGORY_COLORS.items() if k != SAME}
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.map(_clip).to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for i, column in enumerate(frame.columns, 1):
                longest = max([len(str(column))] + [len(str(v)) for v in frame[column].head(300) if v is not None])
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 70)
            if "Category" in frame.columns:
                col = list(frame.columns).index("Category") + 1
                for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                    fill = fills.get(row[0].value)
                    if fill:
                        row[0].fill = fill
            if name == "Failures":
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 6)
    return buffer.getvalue()


def _table(frame: pd.DataFrame, category_col: str | None = None) -> str:
    if frame.empty:
        return "<p class='muted'>None.</p>"
    head = "".join(f"<th>{html.escape(str(c))}</th>" for c in frame.columns)
    names_to_key = {v: k for k, v in CATEGORY_NAMES.items()}
    body = []
    for _, row in frame.iterrows():
        style = ""
        if category_col and row[category_col] in names_to_key:
            style = f" style='background:#{CATEGORY_COLORS[names_to_key[row[category_col]]]}'"
        cells = "".join(f"<td>{'' if pd.isna(v) else html.escape(str(v))}</td>" for v in row)
        body.append(f"<tr{style}>{cells}</tr>")
    return f"<div class='scroll'><table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table></div>"


def html_report(started: datetime, sources: dict[str, dict[str, str]], results: dict[str, RunResult],
                rows: list[Comparison]) -> str:
    cases = cases_frame(rows)
    differences = cases[cases["Category"] != CATEGORY_NAMES[SAME]]
    summary = "".join(f"<tr><th>{html.escape(str(k))}</th><td>{html.escape(str(v))}</td></tr>"
                      for k, v in summary_items(started, sources, results, rows))
    failures = []
    for _, f in failures_frame(results).iterrows():
        failures.append(
            f"<details><summary><b>{html.escape(f['Side'])}</b> · {html.escape(str(f['Case'] or f['Test']))} · "
            f"{html.escape(f['Status'])}</summary><p><b>Test:</b> {html.escape(f['Test'])}</p>"
            f"<p><b>Output:</b> {html.escape(str(f['Output'] or ''))}</p>"
            f"<p><b>Message:</b> {html.escape(f['Message'] or '')}</p>"
            f"<pre>{html.escape(f['Stack trace'] or '')}</pre></details>")
    logs = "".join(f"<details><summary>{SIDE_NAMES[side]} build log</summary><pre>{html.escape(r.log or '')}</pre>"
                   f"</details>" for side, r in results.items())
    counts = category_counts(rows)
    title = f"DiffLens report {started:%Y-%m-%d %H:%M}"
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}</title>
<style>
  body {{ font: 14px/1.45 system-ui, -apple-system, "Segoe UI", sans-serif; color: #0F172A; background: #fff;
         margin: 0 auto; padding: 24px 16px; max-width: 1400px; }}
  h1 {{ font-size: 22px; margin: 0 0 4px; }} h2 {{ font-size: 17px; margin: 28px 0 8px; }}
  .muted {{ color: #64748B; }}
  .cards {{ display: flex; flex-wrap: wrap; gap: 8px; margin: 16px 0; }}
  .card {{ border: 1px solid #E2E8F0; border-radius: 8px; padding: 10px 14px; min-width: 120px; }}
  .card b {{ display: block; font-size: 22px; }}
  .scroll {{ overflow-x: auto; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 12.5px; }}
  th, td {{ border: 1px solid #E2E8F0; padding: 4px 8px; text-align: left; vertical-align: top; }}
  thead th {{ background: #F1F5F9; position: sticky; top: 0; }}
  .summary th {{ width: 260px; background: #F8FAFC; font-weight: 600; }}
  details {{ border: 1px solid #E2E8F0; border-radius: 6px; padding: 6px 10px; margin: 6px 0; }}
  pre {{ white-space: pre-wrap; word-break: break-word; font-size: 12px; background: #F8FAFC; padding: 8px; }}
  @media print {{ details {{ break-inside: avoid; }} details > * {{ display: block; }} }}
</style></head><body>
<h1>{html.escape(title)}</h1>
<p class="muted">Legacy vs new gen, compared input case by input case.</p>
<div class="cards">
  <div class="card"><b>{len(rows)}</b>input cases</div>
  <div class="card"><b>{len(rows) - counts[SAME]}</b>differ</div>
  <div class="card"><b>{counts[REGRESSION]}</b>regressions</div>
  <div class="card"><b>{counts[FIXED]}</b>fixed</div>
  <div class="card"><b>{counts[OUTPUT_DIFF]}</b>output differs</div>
</div>
<h2>Summary</h2><table class="summary">{summary}</table>
<h2>Root causes</h2>
<p class="muted">Words that only one side's output contains, with every differing case they explain.</p>
{_table(causes_frame(rows))}
<h2>Differences ({len(differences)})</h2>{_table(differences, "Category")}
<h2>Failures</h2>{''.join(failures) or "<p class='muted'>No failing tests.</p>"}
<h2>All cases ({len(cases)})</h2>{_table(cases, "Category")}
<h2>Build logs</h2>{logs}
</body></html>"""


def report_name(started: datetime, extension: str) -> str:
    return f"difflens-report-{started:%Y%m%d-%H%M%S}.{extension}"
